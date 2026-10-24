"""The job's city in the resume header (#134)."""

from datetime import date
from pathlib import Path

import pytest
from docx import Document

from src.config import settings
from src.endpoint import relocation as r

HOME = "Kochi, Kerala, India"


@pytest.mark.parametrize("location, city", [
    ("Bengaluru, Karnataka, India", "Bengaluru"),
    ("Bangalore Urban, Karnataka, India", "Bengaluru"),
    ("Bengaluru East, Karnataka, India", "Bengaluru"),
    ("Pune Division, Maharashtra, India", "Pune"),
    ("Gurgaon, Haryana, India", "Gurugram"),
    ("Trivandrum, Kerala, India", "Trivandrum"),
    ("Thiruvananthapuram, Kerala, India", "Trivandrum"),
    ("Navi Mumbai, MH", "Navi Mumbai"),
    ("New Delhi, Delhi, India", "New Delhi"),
])
def test_spelling_variants_resolve_to_one_city(location, city):
    assert r.job_city(location, None, HOME) == city


def test_a_multi_location_job_takes_the_first_city_listed():
    assert r.job_city("Hyderabad / Chennai / Pune", None, HOME) == "Hyderabad"


def test_the_operators_own_city_among_the_jobs_cities_keeps_the_header():
    assert r.job_city("Bengaluru / Kochi", None, HOME) is None
    assert r.job_city("Cochin, Kerala, India", None, HOME) is None


def test_a_remote_role_that_names_a_city_in_the_advert_uses_it():
    assert r.job_city("Remote, IN", "Remote-first, but the team sits in Hyderabad.", HOME) == "Hyderabad"
    assert r.job_city("Remote", "Our Bangalore office hosts a monthly meetup.", HOME) == "Bengaluru"


def test_no_city_anywhere_keeps_the_profile_location():
    assert r.job_city("Remote", "Work from anywhere in India.", HOME) is None
    assert r.job_city(None, None, HOME) is None
    assert r.job_city("", "", HOME) is None


@pytest.mark.parametrize("location", ["KA, IN", "India", "Multiple locations", "Hybrid"])
def test_a_location_without_a_city_falls_back_to_the_advert(location):
    assert r.job_city(location, "Join the team in Pune.", HOME) == "Pune"
    assert r.job_city(location, "No city here.", HOME) is None


def test_a_city_off_the_list_in_the_location_field_is_not_overridden_by_the_advert():
    assert r.job_city("New York, NY", "Offices in New York and Chennai.", HOME) is None


def test_a_city_inside_a_longer_word_does_not_match():
    assert r.cities_in("Punekar Industries, Suratgarh") == []


def test_the_header_line_reads_the_date_unambiguously():
    line = r.header_line("Bengaluru", date(2026, 10, 21), 30)
    assert line == "Relocating to Bengaluru by 20 Nov 2026, willing to relocate sooner if required."
    assert r.header_line("Pune", date(2026, 12, 25), 30).startswith("Relocating to Pune by 24 Jan 2027")


def test_switched_off_the_header_never_changes(monkeypatch):
    monkeypatch.setitem(settings._data, "relocation", {"enabled": False, "days": 30})
    assert r.header_location_for("Pune, MH, India", None) is None


def test_switched_on_a_job_in_another_city_gets_the_line(monkeypatch):
    monkeypatch.setitem(settings._data, "relocation", {"enabled": True, "days": 10})
    monkeypatch.setattr(r, "_home_location", lambda: HOME)
    line = r.header_location_for("Pune, MH, India", None, today=date(2026, 10, 21))
    assert line == "Relocating to Pune by 31 Oct 2026, willing to relocate sooner if required."
    assert r.header_location_for("Kochi, Kerala, India", None) is None


def test_the_relocation_line_is_part_of_the_render_cache_key():
    from src.endpoint.cache import render_version
    from src.llm.schemas import StoredSelection

    sel = StoredSelection(job_id="j", template_version="abc12345", built_at="2026-10-24T00:00:00Z",
                          jd_keywords=[], keyword_coverage=0.0, lead_entry_coverage=0.0, entries=[])
    plain = render_version(sel)
    pune = render_version(sel, "Relocating to Pune by 31 Oct 2026, willing to relocate sooner if required.")
    chennai = render_version(sel, "Relocating to Chennai by 31 Oct 2026, willing to relocate sooner if required.")
    assert len({plain, pune, chennai}) == 3
    assert "_" not in pune, "the cache reads the version back as the key's last _ part"


# --- assembler ---------------------------------------------------------------

REPO = Path(__file__).resolve().parents[1]


def _template() -> Path:
    for p in (REPO / "resumes/templates/headless_v1.docx", REPO / "resumes/templates/example_template.docx"):
        if p.exists():
            return p
    pytest.skip("no Headless template available")


def test_the_assembler_swaps_only_the_location_line(tmp_path):
    import json

    from src.endpoint.assembler import _frozen_canonical, _tailored_start, assemble_docx
    from src.llm.schemas import SelectedEntryOut, StoredSelection

    profile = tmp_path / "p.json"
    profile.write_text(json.dumps({"work_experience": [{"id": "e", "role_blocks": [{
        "role": "data", "entry_header": "Engineer", "bullets": [{"id": "b", "text": "Did a thing."}]}]}],
        "projects": []}))
    sel = StoredSelection(
        job_id="j", template_version="abc12345", built_at="2026-10-24T00:00:00Z",
        jd_keywords=[], keyword_coverage=0.0, lead_entry_coverage=0.0,
        entries=[SelectedEntryOut(kind="work", entry_id="e", block_id="e::data", title_alias="Engineer",
                                  header_left="Engineer", header_right="2025", bullet_ids=["b"], cap=3)],
    )
    line = "Relocating to Pune by 31 Oct 2026, willing to relocate sooner if required."
    plain, moved = tmp_path / "plain.docx", tmp_path / "moved.docx"
    assemble_docx(sel, profile, _template(), plain)
    assemble_docx(sel, profile, _template(), moved, header_location=line)

    a, b = Document(str(plain)), Document(str(moved))
    assert b.paragraphs[2].text == line
    assert a.paragraphs[2].text != line
    for i in (0, 1):
        assert a.paragraphs[i].text == b.paragraphs[i].text
    # Everything else in the frozen prefix is untouched.
    pa = [p for i, p in enumerate(a.element.body) if i != 2]
    pb = [p for i, p in enumerate(b.element.body) if i != 2]
    sa, sb = _tailored_start(a.element.body), _tailored_start(b.element.body)
    assert sa == sb
    from lxml import etree
    assert [etree.tostring(x, method="c14n") for x in pa[:sa - 1]] == \
           [etree.tostring(x, method="c14n") for x in pb[:sb - 1]]
    assert _frozen_canonical(a.element.body, sa) != _frozen_canonical(b.element.body, sb)
