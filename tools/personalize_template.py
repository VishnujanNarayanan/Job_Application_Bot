"""Create your resume template from the example and your master profile.

    python tools/personalize_template.py           # writes the template path in config
    python tools/personalize_template.py --force   # overwrite an existing one

Reads ``resumes/templates/example_template.docx`` (committed, placeholders
only) and fills the header -- name, contact line, location, and the
Education & Certificates lines -- from ``master_profile.yaml``. Optional
display overrides come from ``operator.resume_header`` in config.yaml.

Everything below the Education block is rebuilt by the bot for every job, so
only the header needs your details. The output is gitignored: it carries your
personal information. You can also open it in Word afterwards and adjust the
header by hand; keep the paragraph styles as they are.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "resumes" / "templates" / "example_template.docx"
XMLSPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _phone_display(raw: str) -> str:
    """"+91 98765 43210" for a 10-digit number; anything else as written."""
    digits = re.sub(r"\D", "", raw)
    local, cc = digits[-10:], digits[:-10] or "91"
    return f"+{cc} {local[:5]} {local[5:]}" if len(local) == 10 else raw


def _url_display(url: str) -> str:
    """linkedin.com/in/name: no scheme, no www, no trailing slash."""
    return re.sub(r"^https?://(www\.)?", "", url).rstrip("/")


def _set_text(p, texts: list[str]) -> None:
    """Write ``texts`` into the paragraph's text runs in order; blank the rest."""
    runs = [r for r in p._p.iter(qn("w:r"))
            if r.find(qn("w:t")) is not None and r.find(qn("w:tab")) is None]
    for i, run in enumerate(runs):
        ts = run.findall(qn("w:t"))
        text = texts[i] if i < len(texts) else ""
        ts[0].text = text
        if text != text.strip():
            ts[0].set(XMLSPACE, "preserve")
        for t in ts[1:]:
            t.text = ""


def _fill_contact_line(doc, p, items: list[tuple[str, str]]) -> None:
    """``items`` is (visible text, link) per hyperlink slot, in template order.

    A slot with no value is removed together with its " | " separator, so a
    profile without a LinkedIn or certificates link still reads cleanly.
    """
    links = p._p.findall(qn("w:hyperlink"))
    for link, (text, url) in zip(links, items):
        if not text:
            sep = link.getprevious()
            if sep is None or sep.tag != qn("w:r"):
                sep = link.getnext()
            p._p.remove(link)
            if sep is not None and sep.tag == qn("w:r"):
                p._p.remove(sep)
            continue
        doc.part.rels[link.get(qn("r:id"))]._target = url
        ts = list(link.iter(qn("w:t")))
        ts[0].text = text
        for t in ts[1:]:
            t.text = ""


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--force", action="store_true", help="overwrite an existing template")
    args = parser.parse_args(argv)

    profile_path = ROOT / "master_profile.yaml"
    if not profile_path.exists():
        print("master_profile.yaml not found. Copy master_profile.example.yaml and fill it in first.")
        return 1
    config = yaml.safe_load((ROOT / "config" / "config.yaml").read_text())
    out = ROOT / config["endpoint"]["template_path"]
    if out.exists() and not args.force:
        print(f"{out.relative_to(ROOT)} already exists. Re-run with --force to overwrite it.")
        return 1

    profile = yaml.safe_load(profile_path.read_text())
    me = profile.get("personal") or {}
    header = (config.get("operator") or {}).get("resume_header") or {}

    name = me.get("name") or ""
    if not name:
        print("personal.name is empty in master_profile.yaml.")
        return 1
    phone, email = me.get("phone") or "", me.get("email") or ""
    portfolio = header.get("portfolio_url") or me.get("portfolio") or me.get("github") or ""
    linkedin, certs_link = me.get("linkedin") or "", me.get("certificates_link") or ""

    edu = (profile.get("education") or [{}])[0]
    institution = (edu.get("institution") or "").split("(")[0].strip().rstrip(",")
    edu_line = header.get("education_line") or f"{edu.get('degree', '')}, {institution}".strip(", ")
    edu_status = header.get("education_status") or "Status - Graduated"
    cert_names = [re.sub(r"\s*\([^)]*\)\s*$", "", c.get("name", "")).strip()
                  for c in profile.get("certifications") or []]
    certs_line = header.get("certificates_line") or ", ".join(n for n in cert_names if n)

    doc = Document(str(EXAMPLE))
    paras = doc.paragraphs
    if not paras[4].text.startswith("Education"):
        print("example_template.docx doesn't have the expected layout; was it edited?")
        return 1

    _set_text(paras[0], [name])
    _fill_contact_line(doc, paras[1], [
        (_phone_display(phone) if phone else "", "tel://" + re.sub(r"\D", "", phone)[-10:] + "/"),
        (email, f"mailto:{email}"),
        (_url_display(portfolio) if portfolio else "", portfolio),
        (_url_display(linkedin) if linkedin else "", linkedin),
        ("Certificates" if certs_link else "", certs_link),
    ])
    _set_text(paras[2], [me.get("location") or ""])
    _set_text(paras[5], [edu_line, edu_status])
    _set_text(paras[6], [certs_line])

    doc.core_properties.author = name
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Wrote {out.relative_to(ROOT)} for {name}. Open it in Word to check the header.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
